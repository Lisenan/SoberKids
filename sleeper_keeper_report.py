"""
Sleeper keeper-value report
============================

Pulls last season's FINAL rosters and that season's DRAFT from a Sleeper league,
then shows each manager which round every player on their roster was drafted --
so they can weigh keeper value.

Uses the read-only Sleeper API (no auth / token needed).
Docs: https://docs.sleeper.com/

Usage
-----
    # You have last year's league_id directly:
    python sleeper_keeper_report.py --league-id 123456789012345678

    # You only have THIS year's league_id and want the prior season auto-resolved:
    python sleeper_keeper_report.py --league-id <current_id> --prev

Keeper cost is computed with this league's rules (see keeper_cost() below):
    drafted (non-keeper)  = 1 round earlier      R3    -> R2
    kept last year        = 2 rounds earlier     keptR4 -> R2
    1st round             = stays R1
    waiver / undrafted    = last round - 2       13-rd -> R11

Outputs a console summary, a CSV, and (if pandas+openpyxl are installed) a
formatted Excel workbook grouped by manager.
"""

import argparse
import json
import sys
import time
from pathlib import Path

import requests

try:
    import pandas as pd
    HAVE_PANDAS = True
except ImportError:
    HAVE_PANDAS = False

BASE = "https://api.sleeper.app/v1"
SESSION = requests.Session()
SESSION.headers.update({"User-Agent": "keeper-report/1.0"})

PLAYERS_CACHE = Path("players_nfl.json")
PLAYERS_CACHE_MAX_AGE = 60 * 60 * 24  # refresh at most once/day, per Sleeper docs


# --------------------------------------------------------------------------- #
# HTTP helpers
# --------------------------------------------------------------------------- #
def get(url):
    """GET with basic error handling. Sleeper asks for < 1000 calls/min; we make ~6."""
    resp = SESSION.get(url, timeout=30)
    if resp.status_code == 404:
        return None
    resp.raise_for_status()
    return resp.json()


def load_players():
    """
    The full player map is ~5MB. Cache it locally and only refresh once/day.
    We need it to turn roster player IDs into names -- rosters return IDs only,
    and they include waiver/FA players the draft data never names.
    """
    if PLAYERS_CACHE.exists():
        age = time.time() - PLAYERS_CACHE.stat().st_mtime
        if age < PLAYERS_CACHE_MAX_AGE:
            return json.loads(PLAYERS_CACHE.read_text())
    print("Downloading player map (~5MB, once/day)...", file=sys.stderr)
    data = get(f"{BASE}/players/nfl")
    PLAYERS_CACHE.write_text(json.dumps(data))
    return data


def player_display(pid, players):
    """Human-readable name for a player_id. Handles team D/ST ids like 'DET'."""
    p = players.get(pid)
    if not p:
        return pid  # fall back to the raw id so nothing silently vanishes
    pos = p.get("position")
    first, last = p.get("first_name"), p.get("last_name")
    if pos == "DEF" or not first:
        team = p.get("team") or last or pid
        return f"{team} D/ST"
    return f"{first} {last}".strip()


# --------------------------------------------------------------------------- #
# Data pulls
# --------------------------------------------------------------------------- #
def resolve_league(league_id, use_prev):
    """Return the league object we'll report on (optionally the previous season)."""
    league = get(f"{BASE}/league/{league_id}")
    if league is None:
        sys.exit(f"League {league_id} not found.")
    if use_prev:
        prev = league.get("previous_league_id")
        if not prev:
            sys.exit(f"League {league_id} has no previous_league_id to fall back to.")
        league = get(f"{BASE}/league/{prev}")
        if league is None:
            sys.exit(f"Previous league {prev} not found.")
    return league


def get_managers(league_id):
    """roster_id/owner lookups. Returns user_id -> {'manager', 'team'}."""
    users = get(f"{BASE}/league/{league_id}/users") or []
    out = {}
    for u in users:
        team = (u.get("metadata") or {}).get("team_name") or u.get("display_name")
        out[u["user_id"]] = {"manager": u.get("display_name"), "team": team}
    return out


def get_rosters(league_id):
    return get(f"{BASE}/league/{league_id}/rosters") or []


def get_draft_pick_map(league_id):
    """
    Build player_id -> draft info for the league's draft.

    Keyed by the PLAYER, not by who drafted him: keeper cost almost always
    follows the player's draft slot regardless of later trades. We still record
    who originally drafted him so traded-in players are visible.
    """
    drafts = get(f"{BASE}/league/{league_id}/drafts") or []
    if not drafts:
        return {}, None
    draft = drafts[0]  # sorted most-recent first; most leagues have exactly one
    draft_id = draft["draft_id"]
    picks = get(f"{BASE}/draft/{draft_id}/picks") or []

    pick_map = {}
    for pk in picks:
        pick_map[pk["player_id"]] = {
            "round": pk.get("round"),
            "pick_no": pk.get("pick_no"),
            "is_keeper": bool(pk.get("is_keeper")),
            "drafted_by_user": pk.get("picked_by") or None,
            "drafted_by_roster": pk.get("roster_id"),
        }
    return pick_map, draft


# --------------------------------------------------------------------------- #
# Keeper pricing -- THIS LEAGUE'S RULES
# --------------------------------------------------------------------------- #
def keeper_cost(info, total_rounds):
    """
    Return (keeper_round, basis_str, penalty) for one player.

    League rules (a "penalty" means the keeper costs an EARLIER / more
    expensive pick -- lower round number):

        Drafted, not a keeper last year : 1 round earlier   R3   -> R2
        Kept last year                  : 2 rounds earlier  keptR4 -> R2
        Drafted in round 1              : stays round 1  (no penalty; also the floor)
        Waiver / undrafted              : (last round) - 2   13-rd draft -> R11

    Everything is floored at round 1 -- nothing can cost better than a 1st.
    `penalty` is the EFFECTIVE rounds moved (prior round - keeper round), so a
    floored player shows the real movement rather than the nominal rule.
    """
    # Waiver / free agent: never appears in the draft.
    if info is None or info.get("round") is None:
        if not total_rounds:
            return None, "Waiver (draft length unknown)", None
        kr = max(1, total_rounds - 2)
        return (kr, f"Waiver -> R{kr} (last round {total_rounds} - 2)",
                total_rounds - kr)

    drafted = info["round"]
    kept = bool(info.get("is_keeper"))

    if drafted == 1:
        # No penalty (already the floor) -- but still surface that it was a
        # keeper last year, so managers can see the history at a glance.
        return 1, ("Kept LY R1 -> R1 (no penalty)" if kept
                   else "Drafted R1 -> R1 (no penalty)"), 0

    penalty = 2 if kept else 1
    kr = max(1, drafted - penalty)
    tag = "Kept LY" if kept else "Drafted"
    return kr, f"{tag} R{drafted} -> R{kr} (-{penalty})", drafted - kr


# --------------------------------------------------------------------------- #
# Build the report
# --------------------------------------------------------------------------- #
def build_rows(league_id):
    managers = get_managers(league_id)
    rosters = get_rosters(league_id)
    pick_map, draft = get_draft_pick_map(league_id)
    players = load_players()
    total_rounds = (draft or {}).get("settings", {}).get("rounds")

    rows = []
    for roster in rosters:
        owner_id = roster.get("owner_id")
        mgr = managers.get(owner_id, {"manager": "(orphan/no owner)", "team": None})
        for pid in roster.get("players") or []:
            info = pick_map.get(pid)
            keeper_round, cost_basis, penalty = keeper_cost(info, total_rounds)
            if info and info["round"] is not None:
                drafted_round = info["round"]
                drafted_by = managers.get(info["drafted_by_user"], {}).get("manager")
                status = "Drafted"
                traded_in = drafted_by is not None and drafted_by != mgr["manager"]
                acquired = "Draft"
            else:
                drafted_round = None
                drafted_by = None
                status = "Undrafted (waiver/FA)"
                traded_in = False
                acquired = "Waiver/FA"

            # The round the keeper price is measured FROM: the drafted round, or
            # the last round of the draft for waiver pickups.
            prior_round = drafted_round if drafted_round is not None else total_rounds

            rows.append({
                "manager": mgr["manager"],
                "team_name": mgr["team"],
                "player": player_display(pid, players),
                "position": (players.get(pid) or {}).get("position"),
                "nfl_team": (players.get(pid) or {}).get("team"),
                "kept_last_year": "YES" if (info and info.get("is_keeper")) else "No",
                "acquired": acquired,
                "prior_round": prior_round,     # last year's round (or last rd for waivers)
                "penalty_rounds": penalty,      # rounds moved up
                "keeper_round": keeper_round,   # what it COSTS to keep this year
                "cost_basis": cost_basis,       # plain-English explanation
                "status": status,
                "drafted_round": drafted_round,
                "pick_no": info["pick_no"] if info else None,
                "originally_drafted_by": drafted_by,
                "traded_in": traded_in,
                "player_id": pid,
            })
    return rows, total_rounds


def sort_key(r):
    # group by manager, then by keeper cost (cheapest pick number first = most
    # expensive to keep), unknown-cost rows last.
    return (r["manager"] or "", r["keeper_round"] is None, r["keeper_round"] or 0, r["pick_no"] or 0)


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #
def print_console(rows, league_name=None, season=None, total_rounds=None):
    rows = sorted(rows, key=sort_key)

    # Column widths
    W_KEEP, W_POS, W_PLAYER, W_ORIG, W_NOTE = 5, 4, 24, 8, 26
    width = 2 + W_KEEP + 2 + W_POS + 2 + W_PLAYER + 2 + W_ORIG + 2 + W_NOTE
    bar = "─" * width

    # --- Header -------------------------------------------------------------
    title = "KEEPER VALUE REPORT"
    if league_name:
        title += f" · {league_name}"
    if season:
        title += f" · {season}"
    print()
    print("═" * width)
    print(title.center(width))
    print("═" * width)
    print("  Rules: drafted −1 round · kept last year −2 · 1st round stays R1")
    if total_rounds:
        print(f"  Waivers cost R{max(1, total_rounds - 2)} "
              f"(last round {total_rounds} − 2)   ·   ★ = kept last year")
    else:
        print("  ★ = kept last year")

    # --- Per-manager tables -------------------------------------------------
    current = object()
    for r in rows:
        if r["manager"] != current:
            current = r["manager"]
            label = r["manager"] or "(no owner)"
            if r["team_name"] and r["team_name"] != r["manager"]:
                label += f" — {r['team_name']}"
            print(f"\n {label}")
            print(f" {bar}")
            print(f"  {'KEEP':<{W_KEEP}}  {'POS':<{W_POS}}  {'PLAYER':<{W_PLAYER}}  "
                  f"{'FROM':<{W_ORIG}}  {'BASIS':<{W_NOTE}}")
            print(f" {bar}")

        kr = r["keeper_round"]
        keep_str = f"R{kr}" if kr is not None else "?"
        # ★ marks a player kept last year -- shown even for 1st rounders.
        star = "★" if r["kept_last_year"] == "YES" else " "
        origin = r["originally_drafted_by"] if r["traded_in"] else ""
        name = r["player"][:W_PLAYER]

        print(f"  {keep_str:<{W_KEEP}}{star} {(r['position'] or '—'):<{W_POS}} "
              f"{name:<{W_PLAYER}}  {(origin or '')[:W_ORIG]:<{W_ORIG}}  {r['cost_basis']}")

    # --- Footer summary -----------------------------------------------------
    kept = sum(1 for r in rows if r["kept_last_year"] == "YES")
    waived = sum(1 for r in rows if r["drafted_round"] is None)
    teams = len({r["manager"] for r in rows})
    print(f"\n {bar}")
    print(f"  {len(rows)} players · {teams} teams · "
          f"{kept} kept last year (★) · {waived} from waivers")
    print(f" {bar}")


# Output column order -> friendly header text.
# The four requested columns are the middle block: Keeper?, Prev Round,
# Penalty, and Keeper Round (this year's cost).
COLUMNS = [
    ("manager",               "Manager"),
    ("team_name",             "Team"),
    ("player",                "Player"),
    ("position",              "Pos"),
    ("nfl_team",              "NFL"),
    ("kept_last_year",        "Keeper Last Yr?"),
    ("acquired",              "Acquired"),
    ("prior_round",           "Prev Draft Round"),
    ("penalty_rounds",        "Penalty (Rds)"),
    ("keeper_round",          "Keeper Round 2026"),
    ("cost_basis",            "How It's Calculated"),
    ("pick_no",               "Prev Pick #"),
    ("originally_drafted_by", "Orig. Drafted By"),
    ("traded_in",             "Traded In?"),
    ("player_id",             "Player ID"),
]


def write_outputs(rows, season, csv_path="keeper_report.csv", xlsx_path="keeper_report.xlsx"):
    rows = sorted(rows, key=sort_key)
    keys = [k for k, _ in COLUMNS]
    headers = [h for _, h in COLUMNS]
    # Label the keeper-cost column with the season being played, not last season.
    try:
        headers[keys.index("keeper_round")] = f"Keeper Round {int(season) + 1}"
    except (TypeError, ValueError):
        pass

    if not HAVE_PANDAS:
        import csv
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(headers)
            for r in rows:
                w.writerow([r.get(k) for k in keys])
        print(f"\n  Wrote {csv_path} (pip install pandas openpyxl for Excel)", file=sys.stderr)
        return

    df = pd.DataFrame([[r.get(k) for k in keys] for r in rows], columns=headers)
    df.to_csv(csv_path, index=False)
    print(f"\n  Wrote {csv_path}", file=sys.stderr)

    try:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import CellIsRule

        sheet = f"Keepers {season}"
        with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xl:
            df.to_excel(xl, index=False, sheet_name=sheet, startrow=0)
            ws = xl.sheets[sheet]

            head_fill = PatternFill("solid", fgColor="1F3864")
            head_font = Font(bold=True, color="FFFFFF", size=11)
            thin = Side(style="thin", color="D9D9D9")

            # Header row
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=c)
                cell.fill = head_fill
                cell.font = head_font
                cell.alignment = Alignment(horizontal="center", vertical="center",
                                           wrap_text=True)
            ws.row_dimensions[1].height = 30
            ws.freeze_panes = "D2"           # keep Manager/Team/Player visible
            ws.auto_filter.ref = ws.dimensions

            # Column widths + centering for the short numeric/flag columns
            centered = {"Pos", "NFL", "Keeper Last Yr?", "Acquired",
                        "Prev Draft Round", "Penalty (Rds)", "Prev Pick #",
                        "Traded In?", headers[keys.index("keeper_round")]}
            widths = {"Manager": 16, "Team": 22, "Player": 24, "Pos": 6, "NFL": 6,
                      "Keeper Last Yr?": 14, "Acquired": 11, "Prev Draft Round": 12,
                      "Penalty (Rds)": 11, "How It's Calculated": 32,
                      "Prev Pick #": 11, "Orig. Drafted By": 16,
                      "Traded In?": 10, "Player ID": 12}
            for i, h in enumerate(headers, start=1):
                letter = get_column_letter(i)
                ws.column_dimensions[letter].width = widths.get(h, 16)
                if h in centered:
                    for row in range(2, len(df) + 2):
                        ws.cell(row=row, column=i).alignment = Alignment(horizontal="center")
                for row in range(1, len(df) + 2):
                    ws.cell(row=row, column=i).border = Border(bottom=thin)

            # Highlight last year's keepers in the Keeper? column
            kcol = get_column_letter(keys.index("kept_last_year") + 1)
            ws.conditional_formatting.add(
                f"{kcol}2:{kcol}{len(df) + 1}",
                CellIsRule(operator="equal", formula=['"YES"'],
                           fill=PatternFill("solid", fgColor="FFF2CC"),
                           font=Font(bold=True)))

            # Highlight expensive keepers (rounds 1-3) in the keeper-cost column
            ccol = get_column_letter(keys.index("keeper_round") + 1)
            ws.conditional_formatting.add(
                f"{ccol}2:{ccol}{len(df) + 1}",
                CellIsRule(operator="lessThanOrEqual", formula=["3"],
                           fill=PatternFill("solid", fgColor="FCE4D6"),
                           font=Font(bold=True)))

        print(f"  Wrote {xlsx_path}", file=sys.stderr)
    except Exception as e:  # openpyxl missing / styling issue -- CSV still written
        print(f"  (Skipped Excel: {e})", file=sys.stderr)


# --------------------------------------------------------------------------- #
# Interactive selection: username -> season -> league
# --------------------------------------------------------------------------- #
def get_user(name_or_id):
    """Resolve a Sleeper username (or user_id) to a user object. None if not found."""
    return get(f"{BASE}/user/{name_or_id}")


def default_season():
    """Last completed NFL season, per Sleeper's state endpoint."""
    state = get(f"{BASE}/state/nfl") or {}
    return state.get("previous_season") or state.get("season")


def get_user_leagues(user_id, season):
    return get(f"{BASE}/user/{user_id}/leagues/nfl/{season}") or []


def prompt(msg, default=None):
    suffix = f" [{default}]" if default else ""
    val = input(f"{msg}{suffix}: ").strip()
    return val or (default or "")


def pick_league_interactively(username=None, season=None):
    """Prompt for username, then season, then let the user choose a league."""
    print("\n╔══════════════════════════════════════════╗")
    print("║   SLEEPER KEEPER REPORT                  ║")
    print("╚══════════════════════════════════════════╝")

    # 1. Username -> user_id (re-prompt until found)
    while True:
        uname = username or prompt("  Sleeper username")
        username = None  # only use the passed-in value on the first pass
        if not uname:
            continue
        user = get_user(uname)
        if user and user.get("user_id"):
            break
        print(f"    ✗ No user '{uname}'. Try again.")
    user_id = user["user_id"]
    print(f"    ✓ {user.get('display_name')}")

    # 2. Season -- default to the last completed season, since that's the
    #    one with final rosters + a finished draft for keeper decisions.
    season = season or prompt("  Season", default=default_season())

    # 3. Leagues for that season
    leagues = get_user_leagues(user_id, season)
    if not leagues:
        sys.exit(f"    ✗ No NFL leagues for {user.get('display_name')} in {season}.")

    print(f"\n  Your {season} leagues:")
    for i, lg in enumerate(leagues, 1):
        print(f"    {i}. {(lg.get('name') or '(unnamed)'):<30} "
              f"{lg.get('total_rosters')} teams · {lg.get('status')}")

    if len(leagues) == 1:
        print("    → only one league, using it.")
        return leagues[0]
    print()
    while True:
        raw = prompt(f"  Pick a league [1-{len(leagues)}]", default="1")
        if raw.isdigit() and 1 <= int(raw) <= len(leagues):
            return leagues[int(raw) - 1]
        print("    ✗ Invalid choice.")


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description="Sleeper keeper-value report.")
    ap.add_argument("--league-id",
                    help="Skip the prompts and report on this league_id directly.")
    ap.add_argument("--username",
                    help="Pre-fill the username prompt (interactive mode).")
    ap.add_argument("--season",
                    help="Pre-fill the season prompt. Defaults to the last completed season.")
    ap.add_argument("--prev", action="store_true",
                    help="With --league-id: report on its previous_league_id instead.")
    args = ap.parse_args()

    if args.league_id:
        league = resolve_league(args.league_id, args.prev)
    else:
        league = pick_league_interactively(args.username, args.season)

    league_id, season = league["league_id"], league.get("season")
    league_name = league.get("name")

    rows, total_rounds = build_rows(league_id)
    print_console(rows, league_name=league_name, season=season, total_rounds=total_rounds)
    write_outputs(rows, season)


if __name__ == "__main__":
    main()
